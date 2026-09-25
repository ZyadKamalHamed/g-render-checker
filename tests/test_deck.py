import io
from datetime import datetime

import pytest
from pptx import Presentation

from core import Settings
from core.deck import FOOTER, build_deck
from core.prompt_test import GUARDRAIL, ModelEntry, Prompt, PromptTest, Slot, run_test, summarise
from tests.test_prompt_test import _chain_test


def _texts(slide):
    return " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


@pytest.fixture(scope="module")
def deck_and_test():
    t = _chain_test()
    t.prompts.append(Prompt(id="g", title="Explicit content", kind=GUARDRAIL))
    t.models[0].slots["g"] = Slot(outcome="refused")
    t.recommendation = "Adopt GPT 2.5 Flare for concept renders."
    res = run_test(t, Settings())
    data = build_deck(t, res, summarise(t, res, Settings()), datetime(2026, 9, 25))
    return Presentation(io.BytesIO(data)), t


def test_deck_structure(deck_and_test):
    prs, t = deck_and_test
    # cover, agenda, answer, leaderboard, how, degradation, 3 prompt slides, guardrails, misses,
    # recommendation, appendix, thank you
    assert len(prs.slides) == 14
    assert abs(prs.slide_width / 914400 - 10) < 0.01 and abs(prs.slide_height / 914400 - 5.625) < 0.01
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) >= 2
    assert all(FOOTER in _texts(s) for s in prs.slides)
    assert "Adopt GPT 2.5 Flare" in " ".join(_texts(s) for s in prs.slides)
    assert "Refused" in " ".join(_texts(s) for s in prs.slides)


def test_deck_minimal_test():
    t = _chain_test()
    t.models = t.models[1:]  # one model, no ratings
    for s in t.models[0].slots.values():
        s.rating = None
    res = run_test(t, Settings())
    prs = Presentation(io.BytesIO(build_deck(t, res, summarise(t, res, Settings()), datetime(2026, 9, 25))))
    texts = " ".join(_texts(s) for s in prs.slides)
    assert "Guardrails" not in [s.shapes.title.text if s.shapes.title else "" for s in prs.slides]
    assert "Add your recommendation" in texts
    assert len(prs.slides) >= 10


def test_deck_many_models_and_missing_renders():
    """Nine models (one prompt slide overflows) where most slots are empty or broken."""
    t = _chain_test()
    extra = [ModelEntry(id=f"x{i}", tool="Tool", model=f"Model {i}") for i in range(7)]
    extra[0].slots["p1"] = Slot(image=b"not an image", filename="bad.png")
    t.models += extra
    res = run_test(t, Settings())
    prs = Presentation(io.BytesIO(build_deck(t, res, summarise(t, res, Settings()), datetime(2026, 9, 25))))
    texts = " ".join(_texts(s) for s in prs.slides)
    assert "Not run" in texts
    assert sum("P1 · Realistic" in _texts(s) for s in prs.slides) >= 2  # 9 models → two P1 slides
